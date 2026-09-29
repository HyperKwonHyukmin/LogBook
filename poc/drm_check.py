"""실험 1: 공유 폴더의 문서가 평문으로 읽히는지 확인한다.

사용법: python drm_check.py <파일 또는 폴더> [...] [-o <출력 파일>]
출력: 파일별 {path, size, sniff, open, plaintext} JSON 배열(파일 없으면 stdout, -o 지정 시 그 파일에 UTF-8 로 저장).
종료 코드: 0=DRM·에러 없음, 1=DRM 파일이 하나라도 있음(우선), 3=DRM 은 없지만 stat/open 에러 행이 있음.
"""
import os
import sys
from pathlib import Path

from pocio import dump, parse_args

DRM_MAGIC = b"HHIDRMC"


def sniff(head: bytes) -> str:
    """파일 첫 바이트로 형식을 판별한다."""
    if head.startswith(DRM_MAGIC):
        return "drm"
    if head.startswith(b"PK\x03\x04"):
        return "zip"  # pptx·xlsx·docx 는 zip 컨테이너
    if head.startswith(b"%PDF"):
        return "pdf"
    return "unknown"


def try_open(path: Path) -> str:
    """실제 추출 라이브러리로 열어 본다. 성공 시 'ok ...', 실패 시 'fail ...'."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import pymupdf as fitz  # 'import fitz' 는 deprecated

            with fitz.open(path) as doc:
                if doc.page_count == 0:
                    # DRM 헤더 안에 우연히 %PDF...%%EOF 조각이 섞여 있으면
                    # mupdf 의 복구 모드가 페이지 없는 문서를 "성공"으로 열 수 있다.
                    # 이를 평문으로 오판하지 않도록 실패로 취급한다.
                    return "fail empty/repaired pages=0"
                return f"ok pages={doc.page_count}"
        if ext == ".pptx":
            from pptx import Presentation

            return f"ok slides={len(Presentation(str(path)).slides)}"
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True)
            try:
                return f"ok sheets={len(wb.sheetnames)}"
            finally:
                wb.close()
        return "skip"
    except Exception as exc:  # 실험 도구: 어떤 실패든 기록만 한다
        return f"fail {type(exc).__name__}: {exc}"


def check(path: Path) -> dict:
    """파일 하나를 조사한다. 어떤 파일이 문제여도 전체 실행이 죽지 않게
    stat/open 단계의 OSError 는 잡아서 에러 행으로 남긴다."""
    try:
        with open(path, "rb") as f:
            head = f.read(16)
        size = path.stat().st_size
    except OSError as exc:
        return {
            "path": str(path),
            "size": None,
            "sniff": f"error {type(exc).__name__}: {exc}",
            "open": "skip",
            "plaintext": False,
        }
    sn = sniff(head)
    op = try_open(path)
    return {
        "path": str(path),
        "size": size,
        "sniff": sn,
        "open": op,
        # zip(pptx/xlsx)·pdf 형식인데 실제로 열려야 "평문으로 읽힌다"고 본다.
        "plaintext": sn in ("zip", "pdf") and op.startswith("ok"),
    }


def collect(args: list[str]) -> tuple[list[Path], list[dict]]:
    """대상 파일 목록을 모은다.

    ⚠ 2026-09-28 최종 리뷰: `Path.rglob()` 은 하위 폴더를 나열하다가(권한 없는
    폴더 등) OSError 가 나면 전체 실행을 그대로 끊는다. 다른 실험 스크립트
    (owner_check.collect_owners, bdf_bench.run)와 같은 관례로 `os.walk(onerror=...)`
    를 써서, 순회 자체의 실패는 전체 실행을 죽이지 않고 에러 행으로만 남긴다.

    반환: (파일 Path 목록, 순회 실패를 나타내는 check() 스키마의 에러 행 목록).
    """
    targets: list[Path] = []
    walk_errors: list[dict] = []
    for a in args:
        p = Path(a)
        if not p.is_dir():
            targets.append(p)
            continue

        def on_walk_error(exc: OSError) -> None:
            walk_errors.append({
                "path": getattr(exc, "filename", None) or str(p),
                "size": None,
                "sniff": f"error walk {type(exc).__name__}: {exc}",
                "open": "skip",
                "plaintext": False,
            })

        found: list[Path] = []
        for dirpath, _dirnames, filenames in os.walk(str(p), onerror=on_walk_error):
            for name in filenames:
                found.append(Path(dirpath) / name)
        targets += sorted(found)
    return targets, walk_errors


def main(argv: list[str]) -> int:
    positionals, out = parse_args(argv)
    targets, walk_error_rows = collect(positionals)
    rows = [check(p) for p in targets] + walk_error_rows
    dump(rows, out)
    if any(r["sniff"] == "drm" for r in rows):
        return 1
    if any(r["sniff"].startswith("error") for r in rows):
        return 3
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
