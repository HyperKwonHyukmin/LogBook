"""22 — 보고서 전체 본문 추출. 행·크기를 조용히 자르지 않고, 읽지 못한 부분은 경고로 남겨 '일부 추출'(partial)로 표시한다."""
import io
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

VERSION = "report-fulltext-v2"
SUPPORTED = {".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".ppt"}
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _add(chunks, locator, text):
    text = str(text).strip()
    # 긴 쪽·셀은 글자를 버리지 않고 1800자 창·1544자 간격(256자 겹침)으로 나눈다 — 임베딩 입력 길이 제한용.
    # 검색은 search.keyword_chunks 가 겹침을 빼고 다시 잇는다(256자 이하 검색어는 반드시 한 조각 안에 있다).
    for i in range(0, len(text), 1544):
        chunks.append({"locator": locator if len(text) <= 1800 else f"{locator}:part:{i//1544+1}", "text": text[i:i+1800]})
        if i + 1800 >= len(text):
            break


def tessdata_dir(config):
    """OCR 언어 데이터 폴더: LOGBOOK_REPORT_TESSDATA 가 있으면 그것, 없으면 Tesseract 설치 폴더의 tessdata."""
    if config.tessdata:
        return Path(config.tessdata)
    tool = tesseract_executable()
    return Path(tool).parent / "tessdata" if tool else None


def extraction_identity(config):
    """추출 결과를 가르는 조건(추출기 판·OCR 켜짐·언어·도구·언어 데이터). 바뀌면 캐시 위치가 달라져 다시 추출한다."""
    tool = tesseract_executable()
    signatures = []
    if config.ocr:
        data = tessdata_dir(config)
        for language in config.ocr_languages.split("+"):
            path = (data / f"{language}.traineddata") if data else Path(f"{language}.traineddata")
            try:
                stat = path.stat()
                signatures.append(f"{language}:{stat.st_size}:{stat.st_mtime_ns}")
            except OSError:
                signatures.append(f"{language}:missing")
    return f"{VERSION}|ocr={config.ocr}|langs={config.ocr_languages}|tool={tool}|data={','.join(signatures)}"


def tesseract_executable():
    return shutil.which("tesseract") or (r"C:\Program Files\Tesseract-OCR\tesseract.exe" if Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe").is_file() else None)


def _ocr(data, config):
    executable = tesseract_executable()
    if not config.ocr or not executable:
        raise RuntimeError("이미지 속 글자 미색인: OCR 꺼짐 또는 Tesseract 없음")
    with tempfile.TemporaryDirectory(prefix="logbook-ocr-") as folder:
        source = Path(folder) / "image.png"
        from PIL import Image, UnidentifiedImageError
        try:
            Image.open(io.BytesIO(data)).convert("RGB").save(source)
        except (UnidentifiedImageError, OSError):
            if os.name != "nt" or not data.startswith((b"II\xbc\x01", b"MM\x01\xbc")):
                raise RuntimeError("지원하지 않는 그림 형식")
            # JPEG-XR 은 Windows WIC(HD Photo 코덱)로 PNG 로 바꾼다 — 임시 폴더 안 파일만 다룬다
            encoded = Path(folder) / "image.bin"
            encoded.write_bytes(data)
            decoder_script = Path(folder) / "decode.ps1"
            decoder_script.write_text('''param([string]$SourceFile,[string]$OutputFile)
$ErrorActionPreference = 'Stop'
Add-Type -AssemblyName PresentationCore
$stream = [IO.MemoryStream]::new([IO.File]::ReadAllBytes($SourceFile))
try {
  $decoder = [Windows.Media.Imaging.BitmapDecoder]::Create($stream,[Windows.Media.Imaging.BitmapCreateOptions]::PreservePixelFormat,[Windows.Media.Imaging.BitmapCacheOption]::OnLoad)
  $encoder = [Windows.Media.Imaging.PngBitmapEncoder]::new()
  $encoder.Frames.Add($decoder.Frames[0])
  $output = [IO.File]::Create($OutputFile)
  try { $encoder.Save($output) } finally { $output.Dispose() }
} finally { $stream.Dispose() }
''', encoding="utf-8")
            powershell = shutil.which("powershell.exe") or shutil.which("pwsh.exe")
            if not powershell:
                raise RuntimeError("JPEG-XR 변환 도구(PowerShell) 없음")
            subprocess.run([powershell, "-NoProfile", "-NonInteractive", "-File", str(decoder_script), str(encoded), str(source)],
                           capture_output=True, check=True, timeout=config.conversion_timeout)
        tessdata = tessdata_dir(config)
        args = [executable, str(source), "stdout", "-l", config.ocr_languages]
        if tessdata and tessdata.is_dir():
            args.extend(["--tessdata-dir", str(tessdata)])
        result = subprocess.run(args,
                                capture_output=True, timeout=config.conversion_timeout, check=True)
        return result.stdout.decode("utf-8", errors="replace")


def _office_images(z, chunks, warnings, config):
    images = [n for n in z.namelist() if "/media/" in n and not n.endswith("/")]
    for image in images:
        try:
            _add(chunks, f"image:{image}", _ocr(z.read(image), config))
        except Exception as exc:
            warnings.append(f"{image}: {exc}")


def extract_file(path, config, source_ext=None):
    ext = (source_ext or Path(path).suffix).lower()
    chunks, warnings = [], []
    if ext == ".pdf":
        import fitz
        with fitz.open(stream=Path(path).read_bytes(), filetype="pdf") as document:
            for i, page in enumerate(document, 1):
                text = page.get_text("text", sort=True)
                _add(chunks, f"page:{i}", text)
                if page.get_images() or not text.strip():
                    try:
                        # 쪽 전체를 그림으로 그려 OCR — 그림 속 글자와 윤곽선으로 바뀐 글자까지 잡는다
                        _add(chunks, f"page:{i}:ocr", _ocr(page.get_pixmap(dpi=200).tobytes("png"), config))
                    except Exception as exc:
                        warnings.append(f"page:{i}: {exc}")
    elif ext == ".docx":
        # OOXML 을 직접 읽어 문단·표·글상자·머리말·꼬리말·주석·각주·미주를 모두 담는다
        with zipfile.ZipFile(path) as z:
            parts = [n for n in z.namelist() if n == "word/document.xml" or
                     n.startswith(("word/header", "word/footer", "word/footnotes", "word/endnotes", "word/comments")) and n.endswith(".xml")]
            for part in parts:
                root = ET.fromstring(z.read(part))
                for i, paragraph in enumerate(root.iter(f"{{{NS['w']}}}p"), 1):
                    section = "" if part == "word/document.xml" else Path(part).stem + ":"
                    _add(chunks, f"{section}paragraph:{i}", "".join(t.text or "" for t in paragraph.iter(f"{{{NS['w']}}}t")))
            _office_images(z, chunks, warnings, config)
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        # 읽기 전용 두 벌로 수식과 저장된 계산 값을 함께 담는다(Excel 로 다시 계산하지 않는다)
        formula_stream, value_stream = open(path, "rb"), open(path, "rb")
        formulas = values = None
        try:
            formulas = load_workbook(formula_stream, read_only=True, data_only=False)
            values = load_workbook(value_stream, read_only=True, data_only=True)
            for sheet in formulas:
                sheet.reset_dimensions()
                value_sheet = values[sheet.title]
                value_sheet.reset_dimensions()
                header = ""
                value_rows = iter(value_sheet.iter_rows())
                for row_no, row in enumerate(sheet.iter_rows(), 1):
                    cached_row = next(value_rows, ())
                    cells = []
                    for column, cell in enumerate(row, 1):
                        if cell.value is None:
                            continue
                        coordinate = f"{get_column_letter(column)}{row_no}"
                        value = str(cell.value)
                        cached = cached_row[column-1].value if column <= len(cached_row) else None
                        if cell.data_type == "f":
                            value += f" [cached={cached}]"
                            if cached is None:
                                warnings.append(f"{sheet.title}!{coordinate}: 수식의 저장된 계산 값 없음")
                        cells.append(f"{coordinate}={value}")
                    if cells:
                        if not header:
                            header = " | ".join(cells)[:1000]
                        _add(chunks, f"sheet:{sheet.title}:row:{row_no}", f"Sheet: {sheet.title}\nHeaders: {header}\n" + " | ".join(cells))
        finally:
            if formulas:
                formulas.close()
            if values:
                values.close()
            formula_stream.close()
            value_stream.close()
        with zipfile.ZipFile(path) as z:
            _office_images(z, chunks, warnings, config)
            if any(n.startswith("xl/charts/") and n.endswith(".xml") for n in z.namelist()):
                warnings.append("차트 그림의 글자는 다 읽지 못함(원본 셀은 색인됨)")
            for name in z.namelist():
                if name.startswith("xl/comments") and name.endswith(".xml"):
                    root = ET.fromstring(z.read(name))
                    _add(chunks, name, " ".join(e.text or "" for e in root.iter() if e.tag.endswith("}t")))
    elif ext == ".pptx":
        with zipfile.ZipFile(path) as z:
            # 슬라이드 순서는 presentation.xml 관계 순서를 따른다(파일 이름 순서와 다를 수 있다)
            pns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main", "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
            presentation = ET.fromstring(z.read("ppt/presentation.xml"))
            relations = {r.get("Id"): r.get("Target") for r in ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))}
            import posixpath
            for i, slide in enumerate(presentation.findall(".//p:sldId", pns), 1):
                target = relations[slide.get(f"{{{pns['r']}}}id")]
                part = posixpath.normpath(posixpath.join("ppt", target)) if not target.startswith("/") else target.lstrip("/")
                root = ET.fromstring(z.read(part))
                _add(chunks, f"slide:{i}", "\n".join(t.text or "" for t in root.iter() if t.tag.endswith("}t")))
                relationship_part = posixpath.join(posixpath.dirname(part), "_rels", posixpath.basename(part)+".rels")
                if relationship_part in z.namelist():
                    for relation in ET.fromstring(z.read(relationship_part)):
                        if relation.get("Type", "").endswith("/notesSlide"):
                            notes = posixpath.normpath(posixpath.join(posixpath.dirname(part), relation.get("Target")))
                            note_root = ET.fromstring(z.read(notes))
                            _add(chunks, f"slide:{i}:notes", "\n".join(t.text or "" for t in note_root.iter() if t.tag.endswith("}t")))
                if any(e.tag.endswith("}chart") for e in root.iter()):
                    warnings.append(f"slide:{i}: 차트 그림의 글자는 다 읽지 못함")
            _office_images(z, chunks, warnings, config)
    elif ext == ".ppt":
        with tempfile.TemporaryDirectory(prefix="logbook-ppt-") as folder:
            source = Path(folder) / "source.ppt"
            shutil.copyfile(path, source)  # 변환은 임시 사본으로만 — 원본을 열거나 바꾸지 않는다
            output = Path(folder) / "converted"
            output.mkdir()
            executable = shutil.which("soffice") or shutil.which("libreoffice")
            fallback = Path(r"C:\Program Files\LibreOffice\program\soffice.exe")
            if not executable and fallback.is_file():
                executable = str(fallback)
            if executable:
                subprocess.run([executable, f"-env:UserInstallation={Path(folder,'profile').as_uri()}", "--headless", "--convert-to", "pptx", "--outdir", str(output), str(source)], timeout=config.conversion_timeout, check=True, capture_output=True)
                converted = output / "source.pptx"
            elif os.name == "nt":
                converted = output / "source.pptx"
                # COM 변환은 별도 프로세스로(시간 제한) — 다른 PowerPoint 프로세스를 끝내지 않는다
                script = "import sys,win32com.client; app=win32com.client.DispatchEx('PowerPoint.Application'); p=None\ntry:\n p=app.Presentations.Open(sys.argv[1],True,False,False); p.SaveAs(sys.argv[2],24)\nfinally:\n if p: p.Close()\n app.Quit()"
                import sys
                subprocess.run([sys.executable, "-c", script, str(source), str(converted)], timeout=config.conversion_timeout, check=True, capture_output=True)
            else:
                raise RuntimeError("구형 PPT 는 LibreOffice 또는 Microsoft PowerPoint 변환이 필요함")
            result = extract_file(converted, config)
            result["warnings"].append("구형 PPT 를 임시 사본에서 변환함 — 지원하지 않는 개체는 빠질 수 있음")
            result["state"] = "partial"
            return result
    else:
        raise ValueError(f"지원하지 않는 보고서 형식: {ext}")
    if not chunks:
        warnings.append("검색할 글자를 추출하지 못함")
    warnings = list(dict.fromkeys(warnings))
    return {"chunks": chunks, "warnings": warnings, "state": "partial" if warnings else "ready"}
