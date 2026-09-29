"""PPTX — 슬라이드 본문(표·그룹 포함) + 발표자 노트(설계 §5.2)."""
import io

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .base import ExtractResult, finish_summary

MAX_SLIDE_TITLES = 300


def _shape_texts(shape) -> list[str]:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return [t for s in shape.shapes for t in _shape_texts(s)]
    out: list[str] = []
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        out.append(shape.text_frame.text)
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            out.append(" ".join(c.text for c in row.cells if c.text.strip()))
    return out


def extract_pptx(data: bytes) -> ExtractResult:
    prs = Presentation(io.BytesIO(data))
    r = ExtractResult()
    titles: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        titles.append(title_shape.text_frame.text.strip() if title_shape is not None else "")
        texts = [t for shape in slide.shapes for t in _shape_texts(shape)]
        if not r.add(f"slide:{i}", "\n".join(texts)):
            break
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            if not r.add(f"notes:{i}", slide.notes_slide.notes_text_frame.text):
                break
    cp = prs.core_properties
    finish_summary(r, unit="slide", count=len(prs.slides), title=cp.title, author=cp.author,
                   created=cp.created, headings=titles, slide_titles=titles[:MAX_SLIDE_TITLES])
    return r
