"""XLSX — 시트별 셀 값(계산값 기준, data_only). 미리보기용 앞 N행 표도 여기서 만든다."""
import io
from datetime import date, datetime

from openpyxl import load_workbook

from .base import ExtractResult, finish_summary

MAX_ROWS = 5000
MAX_COLS = 100


def _cell(v) -> str:
    if isinstance(v, float):
        return f"{v:g}"
    if isinstance(v, datetime):
        return v.isoformat(sep=" ", timespec="minutes") if (v.hour or v.minute) else v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _open(data: bytes):
    return load_workbook(io.BytesIO(data), read_only=True, data_only=True)


def extract_xlsx(data: bytes) -> ExtractResult:
    wb = _open(data)
    try:
        r = ExtractResult()
        names = wb.sheetnames
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(max_row=MAX_ROWS, max_col=MAX_COLS, values_only=True):
                vals = [_cell(v) for v in row if v is not None and _cell(v)]
                if vals:
                    lines.append(" ".join(vals))
            if not r.add(f"sheet:{ws.title}", "\n".join(lines)):
                break
        props = wb.properties
        finish_summary(r, unit="sheet", count=len(names), title=props.title, author=props.creator,
                       created=props.created, headings=names)
        return r
    finally:
        wb.close()


def sheet_preview(data: bytes, name: str | None = None, *, max_rows: int = 200, max_cols: int = 50) -> dict:
    """시트 탭 + 앞 max_rows 행 표(설계 §6.3). 뒤쪽 빈 열은 잘라 낸다."""
    wb = _open(data)
    try:
        # 차트 시트(chartsheet)는 셀이 없어 표로 보일 수 없다 — 워크시트 이름만 받는다
        names = [ws.title for ws in wb.worksheets]
        if not names:
            return {"sheets": [], "name": None, "rows": [], "truncated": False}
        target = name if name in names else names[0]
        rows: list[list[str]] = []
        truncated = False
        for i, row in enumerate(wb[target].iter_rows(max_col=max_cols, values_only=True)):
            if i >= max_rows:
                truncated = True
                break
            rows.append(["" if v is None else _cell(v) for v in row])
        width = max((max((j + 1 for j, v in enumerate(r) if v != ""), default=0) for r in rows), default=0)
        rows = [r[:width] for r in rows]
        while rows and not any(rows[-1]):
            rows.pop()
        return {"sheets": names, "name": target, "rows": rows, "truncated": truncated}
    finally:
        wb.close()
