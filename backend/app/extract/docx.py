"""DOCX — 문단(제목 스타일은 목차로) + 표."""
import io

import docx

from .base import ExtractResult, finish_summary

HEADING_STYLES = ("heading", "제목", "title")


def extract_docx(data: bytes) -> ExtractResult:
    d = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    headings: list[str] = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        lines.append(t)
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith(HEADING_STYLES):
            headings.append(t)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" ".join(cells))
    r = ExtractResult()
    r.add("body", "\n".join(lines))
    cp = d.core_properties
    finish_summary(r, unit="paragraph", count=len(lines), title=cp.title, author=cp.author,
                   created=cp.created, headings=headings)
    return r
