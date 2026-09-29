"""추출기 테스트 — 표본 문서는 라이브러리로 메모리에서 만든다(저장소에 파일을 넣지 않는다)."""
import io
from datetime import datetime

import pytest

from app.extract import can_extract, extract_bytes
from app.extract.base import MAX_CHUNK_CHARS, MAX_TOTAL_CHARS, ExtractResult, clean
from app.extract.xlsx import sheet_preview


def _pdf() -> bytes:
    import fitz

    doc = fitz.open()
    for i, body in enumerate(["HULL NO. 9999 Structural Strength Review", "Result summary page"], start=1):
        page = doc.new_page()
        page.insert_text((72, 72), body)
    doc.set_metadata({"title": "Strength Review", "author": "Tester", "creationDate": "D:20260901120000+09'00'"})
    doc.set_toc([[1, "1. Overview", 1], [1, "2. Result", 2]])
    data = doc.tobytes()
    doc.close()
    return data


def _pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "9999 호선 계류 구조 검토"
    s1.placeholders[1].text = "선체 구조 강도 평가"
    s1.notes_slide.notes_text_frame.text = "발표자 노트: 보강재 추가 필요"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "결론"
    prs.core_properties.title = "계류 검토"
    prs.core_properties.author = "홍길동"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "응력"
    ws.append(["부재", "응력(MPa)", None])
    ws.append(["L100x100x10", 123.5, None])
    ws2 = wb.create_sheet("요약")
    ws2["A1"] = "허용응력 초과 없음"
    wb.properties.title = "응력 표"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("1. 개요", level=1)
    d.add_paragraph("호선 9999 의 갑판 구조 검토")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "항목"
    t.rows[0].cells[1].text = "결과"
    d.core_properties.title = "갑판 검토"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_can_extract_by_ext():
    assert can_extract(".pdf") and can_extract(".PPTX") and can_extract(".xlsm") and can_extract(".docx")
    assert not can_extract(".bdf") and not can_extract(".ppt") and not can_extract("")


def test_clean_collapses_whitespace():
    assert clean("  a \t b\r\n\r\n\r\n\r\nc\x00 ") == "a b\n\nc"


def test_result_add_respects_total_limit():
    r = ExtractResult()
    # 조각 하나는 MAX_CHUNK_CHARS 로 잘리므로 여러 조각으로 전체 상한 직전까지 채운다
    left = MAX_TOTAL_CHARS - 5
    i = 0
    while left > 0:
        n = min(left, MAX_CHUNK_CHARS)
        i += 1
        assert r.add(f"page:{i}", "x" * n)
        left -= n
    assert not r.truncated
    r.add("page:last", "y" * 100)
    assert sum(len(t) for _, t in r.chunks) == MAX_TOTAL_CHARS
    assert r.truncated
    assert r.add("page:3", "z") is False


def test_pdf_pages_and_summary():
    r = extract_bytes(".pdf", _pdf())
    assert [loc for loc, _ in r.chunks] == ["page:1", "page:2"]
    assert "HULL NO. 9999" in r.chunks[0][1]
    s = r.summary
    assert s["unit"] == "page" and s["count"] == 2
    assert s["title"] == "Strength Review" and s["author"] == "Tester" and s["created"] == "2026-09-01"
    assert s["headings"] == ["1. Overview", "2. Result"]
    assert s["cover"].startswith("HULL NO. 9999")


def test_pptx_slides_notes_titles():
    r = extract_bytes(".pptx", _pptx())
    locs = [loc for loc, _ in r.chunks]
    assert locs == ["slide:1", "notes:1", "slide:2"]
    assert "선체 구조 강도 평가" in r.chunks[0][1]
    assert "보강재 추가 필요" in r.chunks[1][1]
    s = r.summary
    assert s["unit"] == "slide" and s["count"] == 2
    assert s["title"] == "계류 검토" and s["author"] == "홍길동"
    assert s["headings"] == ["9999 호선 계류 구조 검토", "결론"]
    assert s["slide_titles"] == ["9999 호선 계류 구조 검토", "결론"]


def test_xlsx_sheets():
    r = extract_bytes(".xlsx", _xlsx())
    assert [loc for loc, _ in r.chunks] == ["sheet:응력", "sheet:요약"]
    assert "L100x100x10 123.5" in r.chunks[0][1]
    assert r.summary["unit"] == "sheet" and r.summary["count"] == 2
    assert r.summary["headings"] == ["응력", "요약"] and r.summary["title"] == "응력 표"


def test_docx_body_headings_tables():
    r = extract_bytes(".docx", _docx())
    assert [loc for loc, _ in r.chunks] == ["body"]
    assert "갑판 구조 검토" in r.chunks[0][1] and "항목 결과" in r.chunks[0][1]
    assert r.summary["headings"] == ["1. 개요"] and r.summary["title"] == "갑판 검토"


def test_title_falls_back_to_first_line():
    import fitz

    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "First Line Title\nsecond")
    r = extract_bytes(".pdf", doc.tobytes())
    assert r.summary["title"] == "First Line Title"


def test_broken_bytes_raise():
    with pytest.raises(Exception):
        extract_bytes(".pptx", b"not a zip")


def test_sheet_preview_rows_and_trim():
    p = sheet_preview(_xlsx())
    assert p["sheets"] == ["응력", "요약"] and p["name"] == "응력"
    assert p["rows"] == [["부재", "응력(MPa)"], ["L100x100x10", "123.5"]]
    assert p["truncated"] is False
    assert sheet_preview(_xlsx(), "요약")["rows"] == [["허용응력 초과 없음"]]
    assert sheet_preview(_xlsx(), "없는시트")["name"] == "응력"


def test_sheet_preview_truncates_rows():
    from openpyxl import Workbook

    wb = Workbook()
    for i in range(250):
        wb.active.append([i])
    buf = io.BytesIO()
    wb.save(buf)
    p = sheet_preview(buf.getvalue(), max_rows=200)
    assert len(p["rows"]) == 200 and p["truncated"] is True


def test_sheet_preview_ignores_chartsheet_name():
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    ws = wb.active
    ws.title = "값"
    ws.append([1])
    cs = wb.create_chartsheet("차트")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=1))
    cs.add_chart(chart)
    buf = io.BytesIO()
    wb.save(buf)
    p = sheet_preview(buf.getvalue(), "차트")
    assert p["name"] == "값" and p["sheets"] == ["값"] and p["rows"] == [["1"]]
