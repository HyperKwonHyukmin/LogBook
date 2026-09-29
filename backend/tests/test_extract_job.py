import io
import os

import pytest

from app import jobs, models
from app.entries.locate import FileUnavailable, file_path
from app.extract.job import enqueue_extract, run_extract
from app.extract.store import cache_path
from app.storage.paths import to_long


def _pptx(title="9999 호선 계류 구조 검토", body="선체 구조 강도 평가") -> bytes:
    from pptx import Presentation

    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _write(db, storage, f, data: bytes):
    p = file_path(db, storage, f)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "wb") as fh:
        fh.write(data)


def test_enqueue_only_extractable(db, make_entry_file):
    _e, pdf = make_entry_file(name="a.pdf")
    _e2, bdf = make_entry_file(name="m.bdf", kind="model")
    assert enqueue_extract(db, pdf) is True
    assert enqueue_extract(db, bdf) is False
    db.commit()
    assert db.get(models.FileExtract, pdf.id).state == "queued"
    assert db.get(models.FileExtract, bdf.id) is None
    assert [j.type for j in db.query(models.Job)] == ["extract_file"]


def test_run_extract_stores_texts_summary_and_cache(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pptx", sha="1" * 64)
    _write(db, storage, f, _pptx())
    run_extract(db, storage, f)
    row = db.get(models.FileExtract, f.id)
    assert row.state == "done" and row.error is None and row.chars > 0
    assert row.summary["unit"] == "slide" and row.summary["truncated"] is False
    texts = db.query(models.FileText).filter_by(file_id=f.id).order_by(models.FileText.seq).all()
    assert [t.locator for t in texts] == ["slide:1"] and "강도" in texts[0].text
    assert os.path.exists(to_long(cache_path(storage, "1" * 64)))


def test_run_extract_is_idempotent(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pptx")
    _write(db, storage, f, _pptx())
    run_extract(db, storage, f)
    run_extract(db, storage, f)
    assert db.query(models.FileText).filter_by(file_id=f.id).count() == 1


def test_run_extract_uses_cache_without_reading_file(db, storage, make_entry_file):
    _e, f1 = make_entry_file(name="r.pptx", sha="2" * 64)
    _write(db, storage, f1, _pptx())
    run_extract(db, storage, f1)
    _e2, f2 = make_entry_file(name="copy.pptx", sha="2" * 64)  # 디스크에 없음 — 캐시로만
    run_extract(db, storage, f2)
    assert db.get(models.FileExtract, f2.id).state == "done"


def test_drm_flag_skips(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pdf")
    f.drm_encrypted = True
    db.commit()
    run_extract(db, storage, f)
    row = db.get(models.FileExtract, f.id)
    assert row.state == "skipped" and row.error == "drm"


def test_drm_magic_detected_when_reading(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pdf", sha="3" * 64)
    _write(db, storage, f, b"HHIDRMC" + b"\0" * 100)
    run_extract(db, storage, f)
    assert db.get(models.FileExtract, f.id).error == "drm"
    assert db.get(models.File, f.id).drm_encrypted is True


def test_too_large_skips(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pdf")
    f.size = 300 * 1024 * 1024
    db.commit()
    run_extract(db, storage, f)
    assert db.get(models.FileExtract, f.id).error == "too_large"


def test_broken_file_fails_without_raising(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pptx", sha="4" * 64)
    _write(db, storage, f, b"not a zip at all")
    run_extract(db, storage, f)
    row = db.get(models.FileExtract, f.id)
    assert row.state == "failed" and row.error


def test_missing_file_raises_for_retry(db, storage, make_entry_file):
    _e, f = make_entry_file(name="r.pptx", sha="5" * 64)
    with pytest.raises(OSError):
        run_extract(db, storage, f)


def test_trashed_file_skips(db, storage, make_entry_file):
    _e, f = make_entry_file(status="trashed", name="r.pdf")
    run_extract(db, storage, f)
    assert db.get(models.FileExtract, f.id).error == "trashed"


def test_doc_hull_fills_draft_without_hull(db, storage, make_entry_file):
    e, f = make_entry_file(status="draft", hulls=(), name="r.pptx", sha="6" * 64)
    e.hull_evidence = [{"hull_no": "1234", "score": 1, "reasons": ["이름 1개에 등장"]}]
    version = e.version
    db.commit()
    _write(db, storage, f, _pptx(title="HULL NO. 9999 검토"))
    run_extract(db, storage, f)
    e = db.get(models.Entry, e.id)
    assert [h.hull_no for h in db.query(models.EntryHull).filter_by(entry_id=e.id)] == ["9999"]
    assert e.hull_evidence[0]["hull_no"] == "9999" and e.hull_evidence[0]["score"] == 5
    assert "보고서 본문: r.pptx" in e.hull_evidence[0]["reasons"]
    assert e.version == version


def test_doc_hull_adds_score_to_existing_candidate_but_keeps_hull(db, storage, make_entry_file):
    e, f = make_entry_file(status="draft", hulls=("1234",), name="r.pptx", sha="7" * 64)
    e.hull_evidence = [{"hull_no": "1234", "score": 3, "reasons": ["이름 맨 앞: 1234_x"]},
                       {"hull_no": "9999", "score": 1, "reasons": ["이름 1개에 등장"]}]
    db.commit()
    _write(db, storage, f, _pptx(title="9999 호선 검토"))
    run_extract(db, storage, f)
    e = db.get(models.Entry, e.id)
    assert [h.hull_no for h in db.query(models.EntryHull).filter_by(entry_id=e.id)] == ["1234"]
    top = e.hull_evidence[0]
    assert top["hull_no"] == "9999" and top["score"] == 6


def test_doc_hull_ignored_for_confirmed(db, storage, make_entry_file):
    e, f = make_entry_file(status="confirmed", hulls=("1234",), name="r.pptx", sha="8" * 64)
    _write(db, storage, f, _pptx(title="HULL NO. 9999"))
    run_extract(db, storage, f)
    assert db.get(models.Entry, e.id).hull_evidence is None


def test_process_batch_enqueues_extract(db, storage):
    from app.ingest.process import process_batch

    b = models.Batch(key="20260929-111111-aaaa", source="inbox", original_name="9999_검토", uploader="A100001")
    db.add(b)
    db.commit()
    root = storage.staging / b.key / "9999_검토"
    root.mkdir(parents=True)
    (root / "r.pptx").write_bytes(_pptx())
    (root / "m.bdf").write_bytes(b"GRID,1,,0.,0.,0.\n")
    process_batch(db, storage, b)
    types = sorted(j.type for j in db.query(models.Job))
    assert types == ["extract_file"]
    assert db.query(models.FileExtract).count() == 1


def test_worker_runs_extract_job(db, storage):
    from app.ingest.inbox import InboxWatcher
    from app.worker import run_once

    b = models.Batch(key="20260929-222222-bbbb", source="inbox", original_name="9999_검토", uploader="A100001")
    db.add(b)
    db.flush()
    jobs.enqueue(db, "process_batch", b.id)
    db.commit()
    root = storage.staging / b.key / "9999_검토"
    root.mkdir(parents=True)
    (root / "r.pptx").write_bytes(_pptx())
    stats = run_once(db, storage, InboxWatcher(storage))
    assert stats["failed"] == 0 and stats["processed"] == 2
    f = db.query(models.File).one()
    assert db.get(models.FileExtract, f.id).state == "done"


def test_doc_hull_skips_entry_confirmed_after_load(db, storage, make_entry_file):
    """추출 도중 다른 요청이 초안을 확정했으면(잠근 뒤 다시 보면 draft 가 아님) 호선 근거를 건드리지 않는다."""
    from app.extract.job import _apply_doc_hull

    e, f = make_entry_file(status="draft", hulls=(), name="r.pptx")
    db.get(models.Entry, e.id)  # 이 세션에 draft 상태로 올라와 있다
    from app.database import SessionLocal

    other = SessionLocal()
    try:
        oe = other.get(models.Entry, e.id)
        oe.status = "confirmed"
        other.commit()
    finally:
        other.close()
    _apply_doc_hull(db, f, "9999")
    db.commit()
    db.expire_all()
    assert db.get(models.Entry, e.id).hull_evidence is None
    assert db.query(models.EntryHull).filter_by(entry_id=e.id).count() == 0


def test_doc_hull_regex_ignores_year():
    from app.extract.job import find_doc_hull

    assert find_doc_hull([("slide:1", "호선 2026년 검토")]) is None
    assert find_doc_hull([("slide:1", "호선: 9999")]) == "9999"


def test_trashed_file_keeps_done_texts(db, storage, make_entry_file):
    _e, f = make_entry_file(status="trashed", name="r.pdf")
    db.add(models.FileExtract(file_id=f.id, state="done", chars=2))
    db.add(models.FileText(file_id=f.id, seq=0, locator="page:1", text="본문"))
    db.commit()
    run_extract(db, storage, f)
    assert db.get(models.FileExtract, f.id).state == "done"
    assert db.query(models.FileText).filter_by(file_id=f.id).count() == 1


def test_restore_reenqueues_extract(db, storage, make_user, make_entry_file):
    """휴지통에 있는 동안 건너뛴(trashed) 파일은 복원하면 다시 추출 대기열에 들어간다."""
    from app.entries import service

    u = make_user()
    e, f = make_entry_file(status="confirmed", name="r.pptx")
    _e2, g = make_entry_file(entry=e, name="done.pdf")
    _write(db, storage, f, _pptx())
    _write(db, storage, g, b"%PDF")
    db.add(models.FileExtract(file_id=g.id, state="done"))
    db.commit()
    service.trash(db, storage, db.get(models.Entry, e.id), u)
    run_extract(db, storage, db.get(models.File, f.id))
    assert db.get(models.FileExtract, f.id).error == "trashed"
    db.query(models.Job).delete()
    db.commit()
    service.restore(db, storage, db.get(models.Entry, e.id), u)
    db.expire_all()
    assert [j.target_id for j in db.query(models.Job).filter_by(type="extract_file")] == [f.id]
    assert db.get(models.FileExtract, f.id).state == "queued"
    assert db.get(models.FileExtract, g.id).state == "done"
