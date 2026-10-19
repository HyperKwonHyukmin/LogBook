"""22 — 보고서 본문 색인·검색·근거 요약(보고서 화면 `/api/reports` 로 합침).

- 색인 자체(IndexStore)는 원래 SQLite 라 임시 폴더의 SQLite 를 쓴다. 자료·권한·메타데이터는 logbook_test(MySQL, conftest).
- Ollama·Tesseract 는 모의한다(실제 서버·도구를 부르지 않는다).
"""
import json
import sqlite3
import time
import tracemalloc
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import pytest

from app import models
from app.entries import service
from app.report_index import embeddings
from app.report_index.config import IndexConfig
from app.report_index.extractors import VERSION, _add, extract_file
from app.report_index.search import _keyword, keyword_match, query_terms
from app.report_index.store import SCHEMA_VERSION, IndexStore, normalize, retry_delay
from app.report_index.worker import WorkerLock, index_one, run_cycle
from app.routers.reports import get_report_index, rag_gate

MODEL = "bge-m3:latest:sha256-test"


@pytest.fixture
def config(tmp_path):
    return IndexConfig(index_path=tmp_path / "local" / "index.sqlite3", storage_root=tmp_path / "shared", ocr=False)


@pytest.fixture
def store(tmp_path):
    return IndexStore(tmp_path / "idx" / "index.sqlite3")


def _unit(values):
    v = np.asarray(values, dtype=np.float32)
    return v / np.linalg.norm(v)


def put_doc(store, f, e, chunks, *, model=None, vectors=None, state="ready", sha=None):
    """등록 정보와 같은(현재) 색인 문서를 넣는다. sha 를 주면 등록 SHA 가 다른 '오래된' 색인이 된다."""
    store.upsert_document(
        file_id=f.id, entry_id=e.id, name=f.name, path="원본", sha256=sha or f.sha256, version=VERSION,
        metadata={"registered_sha256": sha or f.sha256, "rel_path": f.rel_path, "vault_rel": e.vault_rel, "entry_id": e.id},
        chunks=[{"locator": loc, "text": text, **({"vector": vectors[i]} if vectors else {})}
                for i, (loc, text) in enumerate(chunks)],
        state=state, embedding_model=model if vectors else None)


# ---- 판정 규칙(순수 함수) ----

def test_terms_across_separate_chunks_and_metadata():
    chunks = [{"text": "받침부 응력", "locator": "p1"}, {"text": "피로 검토", "locator": "p9"}]
    ok, hits, score = keyword_match("크레인 프로젝트", chunks, query_terms("크레인 응력 피로"), "and")
    assert ok and score == 3 and len(hits) == 2
    assert not keyword_match("", chunks, query_terms("응력 피로", "phrase"), "phrase")[0]
    assert keyword_match("", chunks, query_terms("응력 좌굴"), "or")[0]
    assert not keyword_match("13496", [], ["3496"], "and")[0]


def test_long_phrase_across_embedding_split():
    phrase = "unique " + "phrase " * 60
    chunks = []
    _add(chunks, "page:1", "x" * 1500 + phrase + "y" * 1800)
    assert len(chunks) > 1
    ok, hits, _ = keyword_match("", chunks, query_terms(phrase.strip(), "phrase"), "phrase")
    assert ok and hits[0]["locator"] == "page:1"


# ---- 추출기 ----

def test_xlsx_full_content_beyond_old_caps(tmp_path, config):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Loads"
    ws["A1"] = "Load header"
    ws["CW5001"] = "deep UNIQUE_LOAD_TOKEN"
    ws["A2"] = "=1+2"
    path = tmp_path / "report.xlsx"
    wb.save(path)
    result = extract_file(path, config)
    assert any("CW5001=deep UNIQUE_LOAD_TOKEN" in c["text"] and "row:5001" in c["locator"] for c in result["chunks"])
    assert any("=1+2 [cached=None]" in c["text"] for c in result["chunks"])
    assert result["state"] == "partial"


def test_docx_tables_headers_and_body(tmp_path, config):
    from docx import Document
    document = Document()
    document.sections[0].header.paragraphs[0].text = "HEADER_TOKEN"
    document.add_paragraph("BODY_TOKEN")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "TABLE_TOKEN"
    path = tmp_path / "report.docx"
    document.save(path)
    text = " ".join(c["text"] for c in extract_file(path, config)["chunks"])
    assert all(token in text for token in ["HEADER_TOKEN", "BODY_TOKEN", "TABLE_TOKEN"])


def test_pdf_text_and_explicit_image_limitation(tmp_path, config, skip_if_local_drm):
    import io

    import fitz
    from PIL import Image
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((50, 50), "PDF_UNIQUE")
    data = io.BytesIO()
    Image.new("RGB", (20, 20), "white").save(data, format="PNG")
    page.insert_image(fitz.Rect(50, 100, 100, 150), stream=data.getvalue())
    path = tmp_path / "report.pdf"
    pdf.save(path)
    pdf.close()
    result = extract_file(path, config)   # OCR 꺼짐 — 그림 쪽은 경고로 남고 '일부 추출'
    assert any("PDF_UNIQUE" in c["text"] for c in result["chunks"])
    assert result["state"] == "partial" and result["warnings"]


def test_pptx_notes_and_tables(tmp_path, config):
    import zipfile
    path = tmp_path / "report.pptx"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("ppt/presentation.xml", '<p:presentation xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><p:sldIdLst><p:sldId id="256" r:id="rId1"/></p:sldIdLst></p:presentation>')
        z.writestr("ppt/_rels/presentation.xml.rels", '<Relationships><Relationship Id="rId1" Target="slides/slide2.xml"/></Relationships>')
        z.writestr("ppt/slides/slide2.xml", '<slide xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:tbl><a:tr><a:tc><a:t>PPT_TABLE_TOKEN</a:t></a:tc></a:tr></a:tbl></slide>')
        z.writestr("ppt/slides/_rels/slide2.xml.rels", '<Relationships><Relationship Type="x/notesSlide" Target="../notesSlides/notesSlide2.xml"/></Relationships>')
        z.writestr("ppt/notesSlides/notesSlide2.xml", '<notes xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>PPT_NOTES_TOKEN</a:t></notes>')
    text = " ".join(c["text"] for c in extract_file(path, config)["chunks"])
    assert "PPT_TABLE_TOKEN" in text and "PPT_NOTES_TOKEN" in text


def test_tessdata_setting_and_default(tmp_path, monkeypatch):
    from app.report_index import config as cfg
    from app.report_index.extractors import extraction_identity, tessdata_dir
    data = tmp_path / "tess"
    data.mkdir()
    (data / "kor.traineddata").write_bytes(b"k")
    monkeypatch.setenv("LOGBOOK_REPORT_TESSDATA", str(data))
    loaded = cfg.load_config()
    assert tessdata_dir(loaded) == data
    assert "kor:1:" in extraction_identity(loaded) and "eng:missing" in extraction_identity(loaded)
    monkeypatch.delenv("LOGBOOK_REPORT_TESSDATA")
    with patch("app.report_index.extractors.tesseract_executable", return_value=str(tmp_path / "Tess" / "tesseract.exe")):
        assert tessdata_dir(cfg.load_config()) == tmp_path / "Tess" / "tessdata"   # 기본 = 설치 폴더


# ---- 저장소: 스키마·FTS·벡터 ----

def test_fts_candidates_short_terms_and_normalization(store):
    f1, f2 = SimpleNamespace(id=1, name="a", sha256="s1", rel_path="a.pdf"), SimpleNamespace(id=2, name="b", sha256="s2", rel_path="b.pdf")
    e = SimpleNamespace(id=1, vault_rel="2026/E1")
    put_doc(store, f1, e, [("page:1", "계류 하중과 ＡＢＣ 응력 검토")])
    put_doc(store, f2, e, [("page:1", "피로 검토 결과")])
    assert store.fts   # 이 PC 의 SQLite 는 trigram 이 된다
    assert store.keyword_files(normalize("계류 하중")) == {1}       # 3글자 이상 — trigram(공백 포함 부분 일치)
    assert store.keyword_files(normalize("abc")) == {1}            # 전각 → NFKC·casefold
    assert store.keyword_files("검토") == {1, 2}                    # 2글자 — instr
    assert store.keyword_files("응") == {1}                         # 1글자
    assert store.keyword_files("없는말") == set()
    with sqlite3.connect(store.path) as db:
        assert db.execute("PRAGMA user_version").fetchone()[0] == SCHEMA_VERSION


def test_vectors_are_float32_blobs_and_replaced_atomically(store):
    f, e = SimpleNamespace(id=7, name="r", sha256="s", rel_path="r.pdf"), SimpleNamespace(id=7, vault_rel="2026/E7")
    put_doc(store, f, e, [("page:1", "응력"), ("page:2", "피로")], model=MODEL, vectors=[_unit([3, 4]), _unit([1, 0])])
    with sqlite3.connect(store.path) as db:
        rows = db.execute("SELECT dim, typeof(data), length(data) FROM vectors ORDER BY chunk_id").fetchall()
    assert rows == [(2, "blob", 8), (2, "blob", 8)]
    doc = store.get_document(7)
    assert doc["chunk_count"] == doc["vector_count"] == 2 and doc["embedding_model"] == MODEL
    with pytest.raises(KeyError):   # 깨진 조각(위치 없음) — 트랜잭션 전체가 되돌아간다
        store.upsert_document(file_id=7, entry_id=7, name="r", path="x", sha256="other", version=VERSION,
                              chunks=[{"text": "locator 없음"}])
    assert store.get_document(7)["sha256"] == "s" and len(store.chunks([7])) == 2
    assert store.keyword_files("응력") == {7}


def test_migrates_codex_v0_index_keeping_vectors(tmp_path):
    path = tmp_path / "old" / "index.sqlite3"
    path.parent.mkdir()
    with sqlite3.connect(path) as db:   # Codex 초안(스키마 0) 그대로
        db.executescript("""
        CREATE TABLE documents (file_id INTEGER PRIMARY KEY, entry_id INTEGER, name TEXT, path TEXT, sha256 TEXT,
          version TEXT, state TEXT, error TEXT, warnings TEXT, metadata TEXT, size INTEGER, mtime_ns INTEGER,
          indexed_at REAL, checked_at REAL, active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE chunks (id INTEGER PRIMARY KEY, file_id INTEGER NOT NULL, seq INTEGER NOT NULL,
          locator TEXT NOT NULL, text TEXT NOT NULL, vector TEXT, UNIQUE(file_id, seq));
        CREATE INDEX ix_report_chunks_file ON chunks(file_id);
        CREATE TABLE worker_state (key TEXT PRIMARY KEY, value TEXT);
        """)
        meta = json.dumps({"registered_sha256": "s", "embedding_model": MODEL, "embedding_state": "ready"})
        db.execute("INSERT INTO documents VALUES (1,1,'r.pdf','p','s','v','ready',NULL,'[]',?,1,1,1,1,1)", (meta,))
        db.execute("INSERT INTO documents VALUES (2,2,'n.pdf','p','t','v','ready',NULL,'[]','{}',1,1,1,1,1)")
        db.execute("INSERT INTO chunks VALUES (10,1,0,'page:1','계류 하중 검토',?)", (json.dumps([0.6, 0.8]),))
        db.execute("INSERT INTO chunks VALUES (11,2,0,'page:1','벡터 없는 문서',NULL)")
    store = IndexStore(path)
    assert store.keyword_files(normalize("하중 검토")) == {1}
    doc = store.get_document(1)
    assert doc["chunk_count"] == 1 and doc["vector_count"] == 1 and doc["embedding_model"] == MODEL
    best = store.semantic_best(MODEL, _unit([3, 4]), [1, 2], 0.5)
    assert set(best) == {1} and best[1][1] == 10 and best[1][0] == pytest.approx(1.0, abs=1e-3)   # 캐시는 float16(오차 ~1e-3)
    assert store.get_document(2)["vector_count"] == 0
    IndexStore(path)   # 다시 열어도(이미 스키마 2) 그대로
    assert store.keyword_files("문서") == {2}


def test_semantic_cache_invalidated_by_generation_and_streaming(store, monkeypatch):
    f1, f2 = SimpleNamespace(id=1, name="a", sha256="s1", rel_path="a"), SimpleNamespace(id=2, name="b", sha256="s2", rel_path="b")
    e = SimpleNamespace(id=1, vault_rel="v")
    put_doc(store, f1, e, [("page:1", "x")], model=MODEL, vectors=[_unit([1, 0])])
    query = _unit([1, 0])
    assert set(store.semantic_best(MODEL, query, [1, 2], 0.5)) == {1}
    put_doc(store, f2, e, [("page:1", "y")], model=MODEL, vectors=[_unit([1, 0.1])])   # 세대가 바뀐다 → 캐시 다시 읽기
    assert set(store.semantic_best(MODEL, query, [1, 2], 0.5)) == {1, 2}
    assert set(store.semantic_best(MODEL, query, [2], 0.5)) == {2}                       # 허용된 파일만
    assert store.semantic_best("다른 모델", query, [1, 2], 0.5) == {}                     # 모델 digest 가 다르면 섞지 않는다
    monkeypatch.setenv("LOGBOOK_REPORT_VECTOR_CACHE_MB", "0")                              # 상한 0 → 나눠 읽기
    assert set(store.semantic_best(MODEL, query, [1, 2], 0.5)) == {1, 2}
    assert store.semantic_best(MODEL, _unit([0, 1]), [1, 2], 0.5) == {}                    # 임계값 미만


# ---- 워커 ----

def _report(tmp_path, token="old"):
    path = tmp_path / "source.docx"
    path.write_text(token)
    return {"file_id": 1, "entry_id": 2, "name": path.name, "path": str(path), "ext": ".docx",
            "metadata": {"registered_sha256": "dbhash", "rel_path": path.name, "vault_rel": "2026/E2", "entry_id": 2}}


def _extractor(path, config):
    return {"chunks": [{"locator": "paragraph:1", "text": Path(path).read_text()}], "warnings": [], "state": "ready"}


def test_incremental_cache_and_failed_refresh_retains_old_index(tmp_path, config):
    source = _report(tmp_path)
    store = IndexStore(config.index_path)
    assert index_one(store, source, config, extractor=_extractor, embed=False) == "ready"
    assert index_one(store, source, config, extractor=lambda *_: pytest.fail("건너뛰어야 한다"), embed=False) == "unchanged"
    Path(source["path"]).write_text("new content")

    def broken(*_):
        raise ValueError("corrupt")
    assert index_one(store, source, config, extractor=broken, embed=False) == "failed"
    assert store.chunks([1])[0]["text"] == "old"           # 실패해도 이전 색인은 남는다
    assert store.get_document(1)["state"] == "failed"
    assert Path(source["path"]).read_text() == "new content"   # 원본은 바꾸지 않는다


def test_failure_backoff_doubles_and_resets_on_change(tmp_path, config):
    source = _report(tmp_path)
    store = IndexStore(config.index_path)
    calls = []

    def broken(*_):
        calls.append(1)
        raise ValueError("corrupt")
    t0 = time.time()
    assert index_one(store, source, config, extractor=broken, embed=False, now=t0) == "failed"
    doc = store.get_document(1)
    assert doc["fail_count"] == 1 and doc["retry_at"] - doc["checked_at"] == pytest.approx(300, abs=1)
    assert index_one(store, source, config, extractor=broken, embed=False, now=t0 + 10) == "backoff"
    assert len(calls) == 1                                       # 기다리는 동안 원본을 다시 읽지 않는다
    assert index_one(store, source, config, extractor=broken, embed=False, now=doc["retry_at"] + 1) == "failed"
    assert store.get_document(1)["fail_count"] == 2              # 같은 실패 → 간격 두 배
    assert [retry_delay(n, 300, 86400) for n in (1, 2, 3, 10, 20)] == [300, 600, 1200, 86400, 86400]
    Path(source["path"]).write_text("changed!")                  # 파일이 바뀌면 바로 다시
    assert index_one(store, source, config, extractor=_extractor, embed=False, now=t0 + 20) == "ready"
    doc = store.get_document(1)
    assert doc["fail_count"] == 0 and doc["retry_at"] is None


def test_missing_original_backs_off(tmp_path, config):
    source = _report(tmp_path)
    Path(source["path"]).unlink()
    store = IndexStore(config.index_path)
    assert index_one(store, source, config, extractor=_extractor, embed=False) == "missing"
    assert index_one(store, source, config, extractor=_extractor, embed=False) == "backoff"


def test_vectors_filled_later_without_rereading_original(tmp_path, config):
    source = _report(tmp_path, "계류 하중")
    store = IndexStore(config.index_path)
    # 임베딩 서버가 없을 때 — 본문만 색인(상태는 ready 그대로)
    assert index_one(store, source, config, extractor=_extractor, embed=True, model_identity=None) == "ready"
    assert store.get_document(1)["vector_count"] == 0
    with patch("app.report_index.embeddings.httpx.post") as post:
        post.return_value = SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"embeddings": [[3, 4]]})
        outcome = index_one(store, source, config, extractor=lambda *_: pytest.fail("원본을 다시 추출하면 안 된다"),
                            embed=True, model_identity=MODEL)
    assert outcome == "embedded"
    doc = store.get_document(1)
    assert doc["vector_count"] == 1 and doc["embedding_model"] == MODEL and doc["state"] == "ready"


def test_embedding_failure_keeps_text_state_and_backs_off(tmp_path, config):
    source = _report(tmp_path, "계류")
    store = IndexStore(config.index_path)
    with patch("app.report_index.embeddings.httpx.post", side_effect=embeddings.httpx.ConnectError("down")):
        assert index_one(store, source, config, extractor=_extractor, embed=True, model_identity=MODEL) == "ready"
        assert store.get_document(1)["state"] == "ready"   # 본문 상태를 '일부'로 바꾸지 않는다
        assert index_one(store, source, config, extractor=_extractor, embed=True, model_identity=MODEL) == "unchanged"
        doc = store.get_document(1)
        assert doc["embed_fail_count"] == 1 and doc["embed_retry_at"] > time.time()
        assert index_one(store, source, config, extractor=_extractor, embed=True, model_identity=MODEL) == "unchanged"
        assert store.get_document(1)["embed_fail_count"] == 1   # 백오프 동안 다시 부르지 않는다


def test_db_outage_does_not_prune_then_complete_scan_prunes(tmp_path, config):
    source = _report(tmp_path)
    run_cycle(config, source=lambda _: [source], extractor=_extractor, embed=False)

    def outage(_):
        raise ConnectionError("database offline")
    with pytest.raises(ConnectionError):
        run_cycle(config, source=outage)
    store = IndexStore(config.index_path)
    assert len(store.documents()) == 1
    run_cycle(config, source=lambda _: [], embed=False)
    assert store.documents() == []


def test_failed_files_do_not_starve_batch(tmp_path, config):
    """백오프 중인 실패 파일은 배치 상한에 세지 않는다 — 새 보고서가 밀리지 않는다."""
    from dataclasses import replace
    config = replace(config, batch_size=1)
    bad = {**_report(tmp_path), "ext": ".hwp"}
    (tmp_path / "ok.docx").write_text("정상")
    good = {**_report(tmp_path), "file_id": 2, "path": str(tmp_path / "ok.docx"), "name": "ok.docx"}
    assert run_cycle(config, source=lambda _: [bad], extractor=_extractor, embed=False) == {"failed": 1}
    time.sleep(0.01)
    assert index_one(IndexStore(config.index_path), good, config, extractor=_extractor, embed=False) == "ready"
    (tmp_path / "ok.docx").write_text("정상, 고침")   # 실패 파일보다 나중에 확인했지만 바뀐 보고서
    # 오래 확인하지 않은 순서라 실패 파일이 먼저 오지만, 백오프 중이라 배치 1개 자리를 차지하지 않는다
    assert run_cycle(config, source=lambda _: [bad, good], extractor=_extractor, embed=False) == {"backoff": 1, "ready": 1}


def test_worker_lock_excludes_second_process(config):
    with WorkerLock(config.index_path):
        with pytest.raises(RuntimeError):
            with WorkerLock(config.index_path):
                pass


def test_embedding_batches_validated():
    def response(*args, **kwargs):
        assert kwargs["json"]["truncate"] is False and len(kwargs["json"]["input"]) <= 16
        return SimpleNamespace(raise_for_status=lambda: None,
                               json=lambda: {"embeddings": [[3, 4]] * len(kwargs["json"]["input"])})
    with patch("app.report_index.embeddings.httpx.post", side_effect=response) as call:
        out = embeddings.embed_texts(["text"] * 33)
    assert len(out) == 33 and call.call_count == 3 and out[0].dtype == np.float32
    assert out[0].tolist() == pytest.approx([0.6, 0.8])
    bad = SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"embeddings": [[float("nan")]]})
    with patch("app.report_index.embeddings.httpx.post", return_value=bad):
        with pytest.raises(embeddings.EmbeddingUnavailable):
            embeddings.embed_texts(["text"])


def test_identity_cached_60s():
    tags = SimpleNamespace(raise_for_status=lambda: None,
                           json=lambda: {"models": [{"name": embeddings.model_name(), "digest": "abc"}]})
    with patch("app.report_index.embeddings.httpx.get", return_value=tags) as get:
        assert embeddings.embedding_identity() == embeddings.model_name() + ":abc"
        assert embeddings.embedding_identity() == embeddings.model_name() + ":abc"
        assert get.call_count == 1
        embeddings.embedding_identity(max_age=0)   # 워커는 늘 새로 확인한다
        assert get.call_count == 2
    embeddings.clear_identity_cache()
    with patch("app.report_index.embeddings.httpx.get", side_effect=embeddings.httpx.ConnectError("x")) as get:
        for _ in range(3):
            with pytest.raises(embeddings.EmbeddingUnavailable):
                embeddings.embedding_identity()
        assert get.call_count == 1                  # 실패도 60초 동안 다시 쓴다


# ---- API(/api/reports) — logbook_test ----

@pytest.fixture
def world(db, client, make_user, make_entry_file, auth_headers, store):
    """작성자 A100001(LV2), 관리자 Z900009(LV3), 게스트 C300003(LV1). 보고서 3개 + 남의 블록 자료 보고서."""
    make_user("A100001", name="작성자")
    make_user("Z900009", name="관리자", level=3)
    make_user("C300003", name="게스트", level=1)
    crane, crane_r = make_entry_file(title="크레인 검토", hulls=(), name="크레인.pdf")
    moor, moor_r = make_entry_file(title="계류 검토", hulls=("3496",), name="계류.pdf")
    plain, plain_r = make_entry_file(title="일반 검토", hulls=("3497",), name="일반.pdf")
    hidden, hidden_r = make_entry_file(title="비공개", name="비밀.pdf", uploaded_by="B200002")
    hidden.access = "blocked"
    db.commit()
    client.app.dependency_overrides[get_report_index] = lambda: store
    return SimpleNamespace(crane=(crane_r, crane), moor=(moor_r, moor), plain=(plain_r, plain),
                           hidden=(hidden_r, hidden), h=auth_headers("A100001"), admin=auth_headers("Z900009"),
                           guest=auth_headers("C300003"))


def _ids(client, h, query):
    r = client.get("/api/reports?" + query, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def test_search_and_or_phrase_digits_parts_short_terms(client, store, world):
    w, h = world, world.h
    put_doc(store, *w.crane, [("page:1", "받침부 응력 검토"), ("page:9", "피로 수명")])
    put_doc(store, *w.moor, [("page:2", "13496 호선 계류 하중"), ("sheet:하중:row:5", "Sheet: 하중\nA5=계류력")])
    long_text = "x" * 1500 + "경계를 넘는 아주 긴 구문 검색 확인" + "y" * 1800
    chunks = []
    _add(chunks, "page:3", long_text)
    put_doc(store, *w.plain, [(c["locator"], c["text"]) for c in chunks])
    crane_id, moor_id, plain_id = w.crane[0].id, w.moor[0].id, w.plain[0].id
    # AND — 서로 다른 조각과 메타데이터(자료 제목 '크레인 검토') 조합
    r = _ids(client, h, "q=크레인 응력 피로")
    assert [i["file_id"] for i in r["items"]] == [crane_id] and r["keyword_total"] == 1
    item = r["items"][0]
    assert item["match_type"] == "keyword" and item["index_state"] == "ready"
    assert {s["locator"] for s in item["snippets"]} == {"page:1", "page:9"}
    # OR · 2글자 · 1글자
    assert {i["file_id"] for i in _ids(client, h, "q=피로 계류력&mode=or")["items"]} == {crane_id, moor_id}
    assert {i["file_id"] for i in _ids(client, h, "q=하중")["items"]} == {moor_id}
    # 정확 구문
    assert _ids(client, h, "q=응력 검토&mode=phrase")["total"] == 1
    assert _ids(client, h, "q=검토 응력&mode=phrase")["total"] == 0
    # 숫자 경계 — 13496 안의 3496 은 본문으로 맞히지 않는다(계류 자료의 호선 3496 은 메타데이터라 맞는다)
    assert [i["file_id"] for i in _ids(client, h, "q=1349")["items"]] == []
    assert [i["file_id"] for i in _ids(client, h, "q=13496")["items"]] == [moor_id]
    # :part: 경계에 걸린 긴 구문
    r = _ids(client, h, "q=경계를 넘는 아주 긴 구문 검색 확인&mode=phrase")
    assert [i["file_id"] for i in r["items"]] == [plain_id] and r["items"][0]["snippets"][0]["locator"] == "page:3"


def test_keyword_path_never_reads_vectors(client, store, world, monkeypatch):
    put_doc(store, *world.crane, [("page:1", "응력")], model=MODEL, vectors=[_unit([1, 0])])
    seen = []
    original = IndexStore._connect

    from contextlib import contextmanager

    @contextmanager
    def traced(self):
        with original(self) as db:
            db.set_trace_callback(seen.append)
            yield db
    monkeypatch.setattr(IndexStore, "_connect", traced)
    assert _ids(client, world.h, "q=응력")["total"] == 1
    assert _ids(client, world.h, "")["total"] == 3
    assert seen and not [s for s in seen if "vectors" in s.lower()]


def test_unindexed_and_stale_reports_use_existing_search(client, db, store, world):
    w, h = world, world.h
    # 계류: 색인 없음 → 03a 본문(MySQL)으로 찾는다
    db.add(models.FileText(file_id=w.moor[0].id, seq=0, locator="page:4", text="계류 장치 응력 집중"))
    # 일반: 등록 SHA 와 다른 옛 색인 → 쓰지 않고(갱신 대기) 03a 로
    put_doc(store, *w.plain, [("page:1", "옛 본문 고유어")], sha="0" * 64)
    db.add(models.FileText(file_id=w.plain[0].id, seq=0, locator="page:1", text="새 본문 장치"))
    # 크레인: 현재 색인이 있으면 색인이 기준 — 03a 에만 있는 낱말로는 맞지 않는다
    put_doc(store, *w.crane, [("page:1", "색인 본문")])
    db.add(models.FileText(file_id=w.crane[0].id, seq=0, locator="page:1", text="장치"))
    db.commit()
    r = _ids(client, h, "q=장치")
    by = {i["file_id"]: i for i in r["items"]}
    assert set(by) == {w.moor[0].id, w.plain[0].id}
    assert by[w.moor[0].id]["index_state"] == "pending" and by[w.moor[0].id]["snippets"][0]["locator"] == "page:4"
    assert by[w.plain[0].id]["index_state"] == "stale"
    assert _ids(client, h, "q=고유어")["total"] == 0
    assert [i["file_id"] for i in _ids(client, h, "q=장치 응력")["items"]] == [w.moor[0].id]
    # OR — 색인 보고서(크레인 '색인 본문')와 기존 검색 보고서(계류 '응력', 일반 '본문')를 함께
    assert {i["file_id"] for i in _ids(client, h, "q=응력 본문&mode=or")["items"]} == {w.moor[0].id, w.plain[0].id, w.crane[0].id}


def test_acl_is_sql_not_index(client, db, store, world):
    """색인에 들어 있어도 권한은 매 요청 SQL 로 본다 — 공개 단계를 바꾸면 색인을 다시 만들지 않아도 바로 반영된다."""
    w = world
    put_doc(store, *w.hidden, [("page:1", "비밀 응력")])
    put_doc(store, *w.crane, [("page:1", "응력")])
    assert [i["file_id"] for i in _ids(client, w.h, "q=응력")["items"]] == [w.crane[0].id]
    w.hidden[1].access = "normal"
    db.commit()
    assert {i["file_id"] for i in _ids(client, w.h, "q=응력")["items"]} == {w.crane[0].id, w.hidden[0].id}
    assert _ids(client, w.guest, "q=응력")["total"] == 0          # LV1 은 공유 자료·자기 자료만
    w.hidden[1].access = "blocked"
    db.commit()
    assert [i["file_id"] for i in _ids(client, w.h, "q=응력")["items"]] == [w.crane[0].id]


def test_filters_facets_paging_exhaustive(client, db, store, make_entry_file, world):
    w = world
    made = []
    for n in range(81):
        e, f = make_entry_file(title=f"보고서 {n}", hulls=(), name=f"r{n:02d}.pdf")
        if n % 2:
            e.no_hull = True   # 공통 참고자료
        made.append((f, e))
    db.commit()
    for f, e in made:
        put_doc(store, f, e, [("page:1", "stress"), ("page:100", "fatigue")])
    r = _ids(client, w.h, "q=stress fatigue&limit=10&offset=70")
    assert r["total"] == r["keyword_total"] == 81 and len(r["items"]) == 10 and len(r["items"][0]["snippets"]) == 2
    r = _ids(client, w.h, "q=stress fatigue&scope=reference")
    assert r["total"] == 40 and r["scopes"] == {"reference": 40, "unclassified": 41}
    assert r["applied"] == {"scope": ["reference"]}
    assert _ids(client, w.h, "q=stress&scope=hull")["total"] == 0


def test_hybrid_semantic_recommendations_and_outage(client, store, world, monkeypatch):
    w = world
    put_doc(store, *w.crane, [("page:1", "받침부 응력")], model=MODEL, vectors=[_unit([1, 0])])
    put_doc(store, *w.moor, [("page:2", "접안 장치 하중")], model=MODEL, vectors=[_unit([0.9, 0.1])])
    put_doc(store, *w.plain, [("page:1", "관계없는 글")], model=MODEL, vectors=[_unit([0, 1])])
    with patch("app.report_index.embeddings.embedding_identity", return_value=MODEL), \
            patch("app.report_index.embeddings.embed_texts", return_value=[_unit([1, 0])]):
        r = _ids(client, w.h, "q=응력&retrieval=hybrid&sort=relevance")
    assert r["semantic_available"] and r["keyword_total"] == 1 and r["semantic_total"] == 1
    assert [(i["file_id"], i["match_type"]) for i in r["items"]] == [(w.crane[0].id, "keyword"), (w.moor[0].id, "semantic")]
    assert r["items"][1]["snippets"][0]["locator"] == "page:2" and r["semantic_exhaustive"] is False
    monkeypatch.setenv("LOGBOOK_REPORT_SEMANTIC_THRESHOLD", "0.999")
    with patch("app.report_index.embeddings.embedding_identity", return_value=MODEL), \
            patch("app.report_index.embeddings.embed_texts", return_value=[_unit([1, 0])]):
        assert _ids(client, w.h, "q=응력&retrieval=hybrid")["semantic_total"] == 0
    with patch("app.report_index.embeddings.embedding_identity", side_effect=embeddings.EmbeddingUnavailable("embedding_server_unavailable")):
        r = _ids(client, w.h, "q=응력&retrieval=hybrid")
    assert r["keyword_total"] == 1 and r["semantic_available"] is False
    assert r["semantic_error"] == "embedding_server_unavailable"


def test_index_unavailable_keeps_list_and_existing_search(client, db, world):
    client.app.dependency_overrides[get_report_index] = lambda: None
    db.add(models.FileText(file_id=world.moor[0].id, seq=0, locator="page:4", text="계류 장치"))
    db.commit()
    r = _ids(client, world.h, "")
    assert r["total"] == 3 and r["index_available"] is False and r["items"][0]["index_state"] is None
    assert [i["file_id"] for i in _ids(client, world.h, "q=장치")["items"]] == [world.moor[0].id]
    s = client.get("/api/reports/index-status", headers=world.h).json()
    assert s["index_available"] is False and s["pending"] == 3


def test_index_status_counts_dates_and_admin_only_detail(client, store, world):
    w = world
    put_doc(store, *w.crane, [("page:1", "응력")], model=MODEL, vectors=[_unit([1, 0])])
    put_doc(store, *w.moor, [("page:1", "계류")])
    store.record_failure(w.plain[0].id, r"\\server\private\path 오류", state="missing", entry_id=w.plain[1].id,
                         name="일반.pdf", metadata={"registered_sha256": w.plain[0].sha256, "rel_path": w.plain[0].rel_path,
                                                   "vault_rel": w.plain[1].vault_rel, "entry_id": w.plain[1].id})
    tags = SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"models": [{"name": embeddings.model_name(), "digest": "d"}]})
    with patch("app.report_index.embeddings.embedding_identity", return_value=MODEL):
        s = client.get("/api/reports/index-status", headers=w.h).json()
    assert (s["total"], s["ready"], s["failed"], s["missing"], s["embedding_ready"]) == (3, 2, 1, 1, 1)
    from datetime import datetime
    assert datetime.fromisoformat(s["last_updated"]).utcoffset().total_seconds() == 0
    assert "index" not in s and "private" not in json.dumps(s, ensure_ascii=False)
    with patch("app.report_index.embeddings.httpx.get", return_value=tags) as get:
        client.get("/api/reports/index-status", headers=w.h)
        a = client.get("/api/reports/index-status", headers=w.admin).json()
        assert get.call_count == 1   # Ollama 확인은 60초 캐시
    assert a["index"]["documents"] == 3 and a["index"]["vectors"] == 1 and "private" not in json.dumps(a)


def _chat(content):
    return SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"message": {"content": content}})


def test_answer_acl_sources_citations(client, store, world):
    w = world
    assert client.post("/api/reports/answer", headers=w.h, json={"question": "요약", "file_ids": [w.hidden[0].id]}).status_code == 404
    r = client.post("/api/reports/answer", headers=w.h, json={"question": "요약", "file_ids": [w.crane[0].id]})
    assert r.status_code == 422 and r.json()["detail"] == "no_indexed_sources"
    put_doc(store, *w.crane, [("page:3", "크레인 응력 10 MPa")])
    with patch("httpx.post", return_value=_chat("응력은 10 MPa입니다 [S1].")):
        r = client.post("/api/reports/answer", headers=w.h, json={"question": "요약", "file_ids": [w.crane[0].id]}).json()
    assert r["citations"][0]["locator"] == "page:3" and r["citations"][0]["file_id"] == w.crane[0].id
    with patch("httpx.post", return_value=_chat("잘못된 출처 [S99]")):
        r = client.post("/api/reports/answer", headers=w.h, json={"question": "요약", "file_ids": [w.crane[0].id]})
    assert r.status_code == 503 and r.json()["detail"] == "rag_unavailable"


def test_answer_concurrency_limits(client, store, world, monkeypatch):
    w = world
    put_doc(store, *w.crane, [("page:3", "응력")])
    body = {"question": "요약", "file_ids": [w.crane[0].id]}
    rag_gate.enter("A100001")   # 같은 사용자의 요약이 아직 도는 중
    try:
        r = client.post("/api/reports/answer", headers=w.h, json=body)
        assert r.status_code == 429 and r.json()["detail"] == "rag_busy_user"
    finally:
        rag_gate.leave("A100001")
    rag_gate.enter("X1")
    rag_gate.enter("X2")   # 전체 상한 2 가 찼다
    try:
        r = client.post("/api/reports/answer", headers=w.h, json=body)
        assert r.status_code == 429 and r.json()["detail"] == "rag_busy"
        monkeypatch.setenv("LOGBOOK_REPORT_RAG_CONCURRENCY", "3")
        with patch("httpx.post", return_value=_chat("응력 [S1]")):
            assert client.post("/api/reports/answer", headers=w.h, json=body).status_code == 200
    finally:
        rag_gate.leave("X1")
        rag_gate.leave("X2")
    assert rag_gate.users == set()   # 끝나면(실패해도) 자리를 돌려준다


def test_old_report_search_api_removed(client, world):
    assert client.get("/api/report-search", headers=world.h).status_code in (404, 405)


# ---- 규모 ----

def _bulk_index(store, reports, chunks_per, dim, rng):
    """합성 색인을 한 트랜잭션으로 넣는다(upsert 를 수천 번 부르면 커밋 비용만 잰다)."""
    syllables = [chr(c) for c in range(0xAC00, 0xAC00 + 400)]
    with store.write() as db:
        cid = 0
        for fid in range(1, reports + 1):
            meta = json.dumps({"registered_sha256": f"s{fid}", "rel_path": f"{fid}.pdf", "vault_rel": "v", "entry_id": fid})
            db.execute("INSERT INTO documents(file_id, entry_id, name, path, sha256, version, state, warnings, metadata, "
                       "active, chunk_count, vector_count, embedding_model) VALUES (?,?,?,?,?,?,?,?,?,1,?,?,?)",
                       (fid, fid, f"{fid}.pdf", "p", f"s{fid}", VERSION, "ready", "[]", meta, chunks_per, chunks_per, MODEL))
            for seq in range(chunks_per):
                cid += 1
                pick = rng.integers(0, len(syllables), size=500)
                words = [syllables[pick[i]] + syllables[pick[i + 1]] for i in range(0, 500, 2)]
                words[0] = "응력"                       # 모든 보고서에 있는 흔한 2글자
                if fid % 500 == 0 and seq == 1:
                    words[10] = "희귀토큰검사"
                text = " ".join(words)
                db.execute("INSERT INTO chunks(id, file_id, seq, locator, text, norm) VALUES (?,?,?,?,?,?)",
                           (cid, fid, seq, f"page:{seq + 1}", text, normalize(text)))
                vector = rng.standard_normal(dim).astype("<f4")
                vector /= np.linalg.norm(vector)
                db.execute("INSERT INTO vectors VALUES (?,?,?,?,?)", (cid, fid, MODEL, dim, vector.tobytes()))
        store._bump(db)


def test_scale_thousands_of_chunks(tmp_path):
    """합성 2,000 보고서 × 3 조각(6,000 조각, 1024차원) — 키워드는 벡터를 읽지 않고 빠르게, 의미 검색은 한 번의 행렬곱."""
    reports, per, dim = 2000, 3, 1024
    store = IndexStore(tmp_path / "scale" / "index.sqlite3")
    _bulk_index(store, reports, per, dim, np.random.default_rng(22))
    rows = {fid: {"metadata": f"{fid}.pdf", "item": {"name": f"{fid}.pdf"}, "doc": {}} for fid in range(1, reports + 1)}
    indexed = set(rows)

    def keyword(q):
        for r in rows.values():
            r["item"].pop("match_type", None)
        return _keyword(None, None, store, rows, indexed, query_terms(q), "and")

    t = time.perf_counter()
    rare = keyword("희귀토큰검사")
    rare_s = time.perf_counter() - t
    assert sorted(rare) == [500, 1000, 1500, 2000]
    t = time.perf_counter()
    hits = keyword("응력")   # 모든 보고서에 맞는 2글자(instr 로 전체 본문을 훑고 다시 판정 — 가장 무거운 키워드 경우)
    common_s = time.perf_counter() - t
    assert len(hits) == reports
    tracemalloc.start()
    keyword("희귀토큰검사")
    rare_peak = tracemalloc.get_traced_memory()[1]
    tracemalloc.stop()
    query = np.random.default_rng(1).standard_normal(dim).astype(np.float32)
    query /= np.linalg.norm(query)
    t = time.perf_counter()
    store.semantic_best(MODEL, query, range(1, reports + 1), -1.0)   # 처음: 행렬 읽기 + 곱
    first_s = time.perf_counter() - t
    t = time.perf_counter()
    best = store.semantic_best(MODEL, query, range(1, reports + 1), -1.0)   # 캐시
    cached_s = time.perf_counter() - t
    assert len(best) == reports
    print(f"\n[규모] 보고서 {reports} · 조각 {reports * per} · 벡터 {dim}차원 | 희귀어 {rare_s:.3f}s(최대 메모리 "
          f"{rare_peak / 1e6:.1f}MB) · 흔한 2글자 {common_s:.3f}s({len(hits)}건) · 의미 첫 {first_s:.3f}s · 캐시 {cached_s:.3f}s")
    # 느슨한 상한(느린 PC 도 통과하게) — 예전 구조는 요청마다 벡터 JSON 전부(수십 MB)를 읽었다
    assert rare_s < 2.0 and common_s < 10.0 and first_s < 10.0 and cached_s < 2.0
    assert rare_peak < 20 * 1024 * 1024
